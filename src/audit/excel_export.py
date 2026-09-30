"""Excel presentation of the shared report, separate from audit business rules."""

from datetime import date, datetime
from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from src.audit.models import AuditReport
from src.audit.report import (
    cases_table,
    detail_table,
    report_description,
    reports_by_kpi,
    safe_link,
    summary_table,
)

NUMBER_FORMAT = "#,##0.00;[Red](#,##0.00);–"
HEADER_COLOR = "205C93"
INTEGER_COLUMNS = {
    "Index",
    "Invoice days",
    "Days in month",
    "Accounts",
    "Invoice periods",
    "Quantities excluded from kWh total",
}


def _report_information(report: AuditReport) -> pd.DataFrame:
    """Collect scope, calculation notes, warnings, and provenance into one worksheet."""
    records = [
        {"Item": "Scope", "Value": report_description(report)},
        {"Item": "Generated (UTC)", "Value": report.generated_at},

        {
            "Item": "Method",
            "Value": "Allocated kWh = billed kWh / invoice days * days in month. Invoice days count effective dates inclusively. Month-start periods exclude the end; month-end periods exclude the start. Other shared dates belong to the next invoice. Printed dates are retained. Previous invoice pattern is rechecked when the next invoice arrives.",
        },
        {
            "Item": "Rounding",
            "Value": "Values are displayed to two decimals; totals use full precision. Billed kWh repeats across months and must not be summed.",
        },
        {
            "Item": "Coverage",
            "Value": "Coverage checks assume listed accounts may be active during selected months. They do not prove all facilities or invoices are present.",
        },
        {"Item": "Checks needing review", "Value": len(report.cases)},
    ]
    for component, scoped in reports_by_kpi(report):
        records.append({"Item": f"{component} allocated kWh",
                        "Value": scoped.details.allocated_kwh.sum(min_count=1)})
    for warning in report.warnings:
        records.append({"Item": "Input warning", "Value": warning})
    for source in report.sources:
        description = (
            f"Modified UTC: {source['Modified (UTC)']} | SHA256: {source['SHA256']}"
        )
        records.append({"Item": source["File"], "Value": description})
    return pd.DataFrame(records)


def _report_tables(report: AuditReport, include: str) -> list[tuple[str, pd.DataFrame]]:
    tables = []
    for component, scoped in reports_by_kpi(report):
        prefix = component.removeprefix("Purchased ").title()
        legacy = component == "Unclassified"
        tables.append(("Summary" if legacy else f"{prefix} Summary", summary_table(scoped)))
        if include in {"Full report", "Summary and details"}:
            tables.append(("Calculation Details" if legacy else f"{prefix} Details", detail_table(scoped)))
        if include == "Full report":
            tables.append(("Exceptions" if legacy else f"{prefix} Exceptions", cases_table(scoped)))
    tables.append(("Report information", _report_information(report)))
    return tables


def _write_values(sheet: Worksheet, frame: pd.DataFrame) -> None:
    """Write actual values, normalizing pandas missing values to empty Excel cells."""
    sheet.append(list(frame.columns))
    for record in frame.itertuples(index=False, name=None):
        row_values = []
        for value in record:
            if pd.isna(value):
                row_values.append(None)
            else:
                row_values.append(value)
        sheet.append(row_values)


def _format_values(sheet: Worksheet) -> None:
    """Preserve source strings literally; format numbers and dates without coercion."""
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str):
                # Prevent source text beginning with '=' from becoming an Excel formula.
                cell.data_type = "s"
            elif isinstance(cell.value, (int, float)):
                header = sheet.cell(1, cell.column).value
                if header in INTEGER_COLUMNS:
                    cell.number_format = "0"
                else:
                    cell.number_format = NUMBER_FORMAT
            elif isinstance(cell.value, (date, datetime)):
                cell.number_format = "yyyy-mm-dd"


def _header_positions(sheet: Worksheet) -> dict[str, int]:
    positions = {}
    for cell in sheet[1]:
        positions[str(cell.value)] = cell.column
    return positions


def _add_verification_formulas(sheet: Worksheet, data_row_count: int) -> None:
    """Add recalculation and difference formulas using named source-column positions."""
    headers = _header_positions(sheet)
    billed_column = get_column_letter(headers["Billed kWh"])
    billing_days_column = get_column_letter(headers["Invoice days"])
    month_days_column = get_column_letter(headers["Days in month"])
    allocated_column = get_column_letter(headers["Allocated kWh"])
    recalculated_index = sheet.max_column + 1
    difference_index = recalculated_index + 1
    recalculated_column = get_column_letter(recalculated_index)
    sheet.cell(1, recalculated_index, "Recalculated kWh")
    sheet.cell(1, difference_index, "Difference kWh")
    for row_number in range(2, data_row_count + 2):
        billed_cell = f"{billed_column}{row_number}"
        billing_days_cell = f"{billing_days_column}{row_number}"
        month_days_cell = f"{month_days_column}{row_number}"
        allocated_cell = f"{allocated_column}{row_number}"
        recalculated_cell = f"{recalculated_column}{row_number}"
        recalculation_formula = (
            f'=IF(OR({billed_cell}="",{billing_days_cell}<=0,{month_days_cell}=""),"",'
            f"{billed_cell}/{billing_days_cell}*{month_days_cell})"
        )
        difference_formula = (
            f'=IF(OR({recalculated_cell}="",{allocated_cell}=""),"",'
            f"{allocated_cell}-{recalculated_cell})"
        )
        sheet.cell(row_number, recalculated_index, recalculation_formula)
        sheet.cell(row_number, difference_index, difference_formula)
        sheet.cell(row_number, recalculated_index).number_format = NUMBER_FORMAT
        sheet.cell(row_number, difference_index).number_format = NUMBER_FORMAT


def _add_source_links(sheet: Worksheet) -> None:
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            header = str(sheet.cell(1, cell.column).value)
            is_source_column = header == "Source"
            if header.startswith("Source "):
                is_source_column = True
            if not is_source_column:
                continue
            if safe_link(cell.value):
                cell.hyperlink = cell.value
                cell.value = "Open evidence file"
                cell.font = Font(color=HEADER_COLOR, underline="single")


def _freeze_columns(sheet: Worksheet) -> None:
    sheet.freeze_panes = "C2"
    if "Unit name" not in _header_positions(sheet) or sheet.title.endswith("Exceptions"):
        return
    headers = _header_positions(sheet)
    unit_column = headers["Unit name"]
    # Excel freezes everything above/left of the specified cell.
    first_scrolling_column = get_column_letter(unit_column + 1)
    sheet.freeze_panes = f"{first_scrolling_column}2"


def _column_width(header: str) -> int:
    """Reserve space for narrative evidence while keeping dates and counts narrow."""
    if header in {"Explanation", "Next step", "Date adjustment", "Reason"}:
        return 60
    if header == "Source" or header.startswith("Source "):
        return 26
    if header in {"Invoice start", "Invoice end", "Allocation start", "Allocation end", "Reporting month"}:
        return 16
    if header in {"Index", "Invoice days", "Accounts", "Days in month"}:
        return 12
    if "kWh" in header:
        return 18
    return 24


def _format_sheet(sheet: Worksheet) -> None:
    """Apply shared print, navigation, header, and wrapping settings."""
    _freeze_columns(sheet)
    sheet.auto_filter.ref = sheet.dimensions
    sheet.print_title_rows = "1:1"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    for header_cell in sheet[1]:
        header_cell.fill = PatternFill("solid", fgColor=HEADER_COLOR)
        header_cell.font = Font(color="FFFFFF", bold=True)
        header_cell.alignment = Alignment(wrap_text=True, vertical="top")
        column_width = _column_width(str(header_cell.value))
        sheet.column_dimensions[header_cell.column_letter].width = column_width
    sheet.row_dimensions[1].height = 32
    # Leave individual row heights automatic so wrapped evidence is never clipped.
    sheet.sheet_format.defaultRowHeight = 15
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    if sheet.title == "Report information":
        sheet.column_dimensions["B"].width = 110


def excel_bytes(report: AuditReport, include: str = "Full report") -> bytes:
    """Select tables, write values/formulas, apply formatting, then serialize in memory."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name, frame in _report_tables(report, include):
        sheet = workbook.create_sheet(sheet_name)
        _write_values(sheet, frame)
        _format_values(sheet)
        if "Billed kWh" in frame.columns:
            _add_verification_formulas(sheet, len(frame))
        _add_source_links(sheet)
        _format_sheet(sheet)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
