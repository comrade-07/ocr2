from datetime import date

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.output.xlsx_writer import write_xlsx
from src.pipeline.run_gold_pipeline import (
    build_gold_template,
    build_gold_template_aggregated,
    gold_output_filename,
    run_gold_pipeline,
)
from src.pipeline.run_silver_pipeline import silver_output_filename


def _as_date(value):
    return value.date() if hasattr(value, "date") else value


def test_build_gold_template_applies_unit_conversions():
    template_output_df = pd.DataFrame([
        {
            "Energy KPI": "Fossil Fuels",
            "Amount of energy consumed": 1.25,
            "Energy Unit": "MWh",
            "Total amount of energy consumed": 2,
            "source_files": "one.json",
        },
        {
            "Energy KPI": "Renewable sources",
            "Amount of energy consumed": 50,
            "Energy Unit": "kWh",
            "Total amount of energy consumed": 100,
            "source_files": "two.json",
        },
    ])

    result = build_gold_template(template_output_df)

    assert list(result.columns) == [
        "Energy KPI",
        "Amount of energy consumed",
        "Energy Unit",
        "Total amount of energy consumed",
        "Original Energy Unit",
        "source_files",
    ]
    assert list(result["Amount of energy consumed"]) == [1250, 50]
    assert list(result["Energy Unit"]) == ["kWh", "kWh"]
    assert list(result["Original Energy Unit"]) == ["MWh", "kWh"]
    assert list(result["Total amount of energy consumed"]) == [2000, 100]
    assert list(result["source_files"]) == ["one.json", "two.json"]


def test_build_gold_template_aggregated_sums_facility_month_rows():
    gold_template_df = pd.DataFrame([
        {
            "Division": "Division A",
            "Legal Entity Name": "Entity A",
            "Unit": "Unit A",
            "Consumption start date": "2026-01-01",
            "Amount of energy consumed": 6047,
            "Energy KPI": "Nuclear sources",
            "source_files": "invoice.json",
            "manual_data_entry_portion": "",
        },
        {
            "Division": "Division A",
            "Legal Entity Name": "Entity A",
            "Unit": "Unit A",
            "Consumption start date": "2026-01-01",
            "Amount of energy consumed": 92.71,
            "Energy KPI": "Nuclear sources",
            "source_files": "manual.msg",
            "manual_data_entry_portion": "Yes",
        },
        {
            "Division": "Division A",
            "Legal Entity Name": "Entity A",
            "Unit": "Unit A",
            "Consumption start date": "2025-12-01",
            "Amount of energy consumed": 89.5,
            "Energy KPI": "Nuclear sources",
            "source_files": "manual.msg",
            "manual_data_entry_portion": "Yes",
        },
    ])

    gold_template_df["KPI Component"] = "Purchased electricity"
    result = build_gold_template_aggregated(gold_template_df)

    assert len(result) == 2
    january_row = result[result["Consumption start date"].astype(str) == "2026-01-01"].iloc[0]
    assert january_row["Amount of energy consumed"] == 6139.71
    assert january_row["source_files"] == "invoice.json; manual.msg"
    assert january_row["manual_data_entry_portion"] == "Yes"


def test_run_gold_pipeline_formats_template_dates_like_silver_step_10(tmp_path):
    config_dir = tmp_path / "config"
    silver_dir = tmp_path / "silver"
    gold_dir = tmp_path / "gold"
    config_dir.mkdir()
    silver_dir.mkdir()
    (config_dir / "settings.yaml").write_text(
        "\n".join([
            "paths:",
            f"  silver_excel_output: {silver_dir.as_posix()}",
            f"  gold_output: {gold_dir.as_posix()}",
            "",
        ]),
        encoding="utf-8",
    )
    silver_template_output = pd.DataFrame([
        {
            "Division": "Division A",
            "Legal Entity Name": "Entity A",
            "Unit": "Unit A",
            "Consumption start date": "2025-06-01",
            "Consumption end date": "2025-06-30",
            "Transaction Date": "2026-06-14",
            "Energy KPI": "Fossil Fuels",
            "Amount of energy consumed": 1.25,
            "Energy Unit": "MWh",
            "Total amount of energy consumed": 2,
        },
        {
            "Division": "Division A",
            "Legal Entity Name": "Entity A",
            "Unit": "Unit A",
            "Consumption start date": "2025-06-01",
            "Consumption end date": "2025-06-30",
            "Transaction Date": "2026-06-14",
            "Energy KPI": "Fossil Fuels",
            "Amount of energy consumed": 0.75,
            "Energy Unit": "MWh",
            "Total amount of energy consumed": 1,
        }
    ])
    silver_template_output["KPI Component"] = "Purchased electricity"
    write_xlsx(
        silver_template_output,
        silver_dir,
        silver_output_filename("scope2", "template_output"),
        sheet_name="SilverTemplateOutput",
    )

    output = run_gold_pipeline(config_dir=config_dir)
    aggregated_output = gold_dir / gold_output_filename("scope2", "template_aggregated")

    workbook = load_workbook(output)
    worksheet = workbook["GoldTemplate"]
    headers = [cell.value for cell in worksheet[1]]

    for header, expected_date in {
        "Consumption start date": date(2025, 6, 1),
        "Consumption end date": date(2025, 6, 30),
        "Transaction Date": date(2026, 6, 14),
    }.items():
        cell = worksheet.cell(row=2, column=headers.index(header) + 1)
        assert _as_date(cell.value) == expected_date
        assert cell.number_format == "yyyy-mm-dd"

    aggregated_workbook = load_workbook(aggregated_output)
    aggregated_worksheet = aggregated_workbook["GoldTemplateAggregated"]
    aggregated_headers = [cell.value for cell in aggregated_worksheet[1]]
    assert aggregated_worksheet.max_row == 2
    assert aggregated_worksheet.cell(
        row=2,
        column=aggregated_headers.index("Amount of energy consumed") + 1,
    ).value == 2000


def test_gold_separates_components_combines_units_and_does_not_count_solar_as_purchases():
    common = {
        "Division": "Division A", "Legal Entity Name": "Entity A", "Unit": "Facility A",
        "Consumption start date": "2026-01-01", "Energy KPI": "Fossil Fuels",
    }
    rows = [
        {**common, "KPI Component": "Purchased electricity", "Energy Unit": "MWh", "Amount of energy consumed": 2, "Total amount of energy consumed": 2, "source_files": "electricity.json"},
        {**common, "KPI Component": "Purchased electricity", "Energy Unit": "kWh", "Amount of energy consumed": 500, "Total amount of energy consumed": 500, "source_files": "electricity.pdf"},
        {**common, "KPI Component": "Purchased electricity", "Energy Unit": "MWh", "Energy KPI": "Renewable energy production", "Amount of energy consumed": 0.2, "Total amount of energy consumed": 2, "source_files": "electricity.json"},
        {**common, "KPI Component": "Purchased steam", "Energy Unit": "ton", "Amount of energy consumed": 10, "Total amount of energy consumed": 10, "source_files": "steam.pdf"},
        {**common, "KPI Component": "Purchased heat", "Energy Unit": "MWh", "Amount of energy consumed": 3, "Total amount of energy consumed": 3, "source_files": "heat.json"},
        {**common, "KPI Component": "Purchased cool", "Energy Unit": "RTH", "Amount of energy consumed": 100, "Total amount of energy consumed": 100, "source_files": "cool.pdf"},
    ]
    gold = build_gold_template(pd.DataFrame(rows))
    repeated = build_gold_template(gold)
    pd.testing.assert_frame_equal(gold, repeated)
    aggregated = build_gold_template_aggregated(gold).set_index("KPI Component")
    expected = {"Purchased electricity": 2500, "Purchased steam": 7200, "Purchased heat": 3000, "Purchased cool": 62}
    assert aggregated["Amount of energy consumed"].to_dict() == expected
    assert aggregated["Total amount of energy consumed"].to_dict() == expected
    assert aggregated.loc["Purchased electricity", "source_files"] == "electricity.json; electricity.pdf"
    assert set(aggregated["Energy Unit"]) == {"kWh"}


def test_gold_rejects_unconverted_aggregation_and_flagged_template_rows():
    rows = pd.DataFrame([{
        "Division": "A", "Legal Entity Name": "B", "Unit": "C",
        "Consumption start date": "2026-01-01", "KPI Component": "Purchased steam",
        "Energy Unit": "ton", "Amount of energy consumed": 10,
    }])
    assert build_gold_template_aggregated(rows).empty
    rows["business_mapping_issue"] = "Energy source allocation: missing allocation percentages"
    assert build_gold_template(rows).empty
    valid = rows.assign(business_mapping_issue="", **{"Energy Unit": "kWh"})
    invalid_unit = valid.assign(**{"Energy Unit": "mmBTU"})
    result = build_gold_template(pd.concat([rows, valid, invalid_unit], ignore_index=True))
    assert len(result) == 1
    assert result.iloc[0]["Amount of energy consumed"] == 10


def test_aggregated_totals_weight_allocations_and_count_source_quantities_once():
    from src.pipeline.run_silver_pipeline import build_silver_template_preparation, build_silver_template_output

    common = {
        "Division": "A", "Legal Entity Name": "B", "Unit": "C", "Energy Type": "Electricity",
        "KPI Component": "Purchased electricity", "Consumption start date": "2026-01-01",
        "nuclear_%": 0, "source_file_count": 1, "proration_component_count": 1,
    }
    business = pd.DataFrame([
        {**common, "Account number": "001", "Energy Unit": "MWh", "Amount of energy consumed": 1,
         "fossil_fuel_%": 0.5, "renewable_energy_%": 0.5, "source_files": "one.pdf",
         "solar_total_generation": 0.1, "total_solar_consumed": 0.05},
        {**common, "Account number": "002", "Energy Unit": "kWh", "Amount of energy consumed": 1000,
         "fossil_fuel_%": 0.2, "renewable_energy_%": 0.8, "source_files": "two.json",
         "solar_total_generation": 100, "total_solar_consumed": 50},
    ])
    gold = build_gold_template(build_silver_template_output(build_silver_template_preparation(business)))
    row = build_gold_template_aggregated(gold).iloc[0]
    assert row["Amount of energy consumed"] == row["Total amount of energy consumed"] == 2000
    assert row["solar_total_generation"] == 200
    assert row["total_solar_consumed"] == 100
    assert row["fossil_fuel_%"] == pytest.approx(0.35)
    assert row["renewable_energy_%"] == pytest.approx(0.65)
    assert row["nuclear_%"] == 0
    assert row["source_file_count"] == row["proration_component_count"] == 2
